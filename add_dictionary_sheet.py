"""Thêm/cập nhật sheet 'Dictionary' trong workbook từ dict/dictionary.json.
Mỗi mục 2 dòng: Dict|<no>|title (tên mục) và Dict|<no>|text (giải thích).
Tên tiếng Anh lấy từ sheet Dic_terms (Dic|<no>); phần giải thích để trống để dịch.
Dùng: python add_dictionary_sheet.py [--xlsx Miramisu Eng.xlsx]
"""
import argparse, json, os, re, shutil, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
HEAD = ['ID', 'Tiếng Nhật / gốc', 'Bản dịch tiếng Anh', 'Loại text', 'Ngữ cảnh', 'Ghi chú của người dịch']
F_HEAD = Font(name='Arial', size=11, bold=True, color='FFFFFF')
F_BODY = Font(name='Arial', size=11, color='252525')
FILL_HEAD = PatternFill('solid', fgColor='314F4C')
FILL_TR = PatternFill('solid', fgColor='FFF3CF')
WRAP = Alignment(wrap_text=True, vertical='top')
BAD = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean(v):
    return BAD.sub('', v) if isinstance(v, str) else v


def main():
    cfg = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=cfg['workbook'])
    ap.add_argument('--dict', default=os.path.join(HERE, 'dict', 'dictionary.json'))
    a = ap.parse_args()
    xlsx = os.path.join(HERE, a.xlsx) if not os.path.isabs(a.xlsx) else a.xlsx
    data = json.load(open(a.dict, encoding='utf-8'))['data']

    print('đang mở', xlsx)
    wb = openpyxl.load_workbook(xlsx)
    # tên tiếng Anh của các mục, lấy từ Dic_terms
    en_title = {}
    if 'Dic_terms' in wb.sheetnames:
        for r in wb['Dic_terms'].iter_rows(min_row=2, values_only=True):
            if r and r[0] and str(r[0]).startswith('Dic|') and r[2]:
                en_title[str(r[0])[4:]] = str(r[2])

    if 'Dictionary' in wb.sheetnames:
        wb.remove(wb['Dictionary'])
    ws = wb.create_sheet('Dictionary')
    ws.freeze_panes = 'A2'
    for i, w in enumerate((22, 46, 46, 22, 24, 34), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, h in enumerate(HEAD, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
    row = 2
    n_title = n_text = 0
    for e in data:
        no, cat = str(e.get('no', '')), str(e.get('category', ''))
        ctx = 'no=%s / %s' % (no, cat)
        t_en = en_title.get(no, '')
        vals = [['Dict|%s|title' % no, e.get('title', ''), t_en or None, 'Từ điển - tên', ctx,
                 'Tên mục (lấy từ Dic_terms)' if t_en else 'Cần dịch tên'],
                ['Dict|%s|text' % no, e.get('text', ''), None, 'Từ điển - giải thích', ctx, 'Cần dịch giải thích']]
        for v in vals:
            for i, x in enumerate(v, 1):
                c = ws.cell(row, i, clean(x))
                c.font, c.alignment = F_BODY, WRAP
                if i == 3:
                    c.fill = FILL_TR
            row += 1
            n_title += (v[3].endswith('tên'))
            n_text += (v[3].endswith('giải thích'))
    # đặt sau Dic_terms
    if 'Dic_terms' in wb.sheetnames:
        wb.move_sheet(ws, offset=wb.sheetnames.index('Dic_terms') + 1 - wb.sheetnames.index('Dictionary'))

    # TOC: thêm dòng sau dòng sheet cuối, cập nhật tổng
    toc = wb['TOC']
    last = 0
    for r in range(7, 400):
        v = toc.cell(r, 1).value
        if v not in (None, ''):
            last = r
        elif last and r > last + 2:
            break
    new_row = next((r for r in range(7, last + 1) if str(toc.cell(r, 1).value) == 'Dictionary'), last + 1)
    vals = ['Dictionary', row - 2, '=B%d-COUNTBLANK(Dictionary!C2:C%d)' % (new_row, row - 1),
            '=C%d/B%d' % (new_row, new_row), 'Từ điển trong game: tên mục + giải thích (Dictionary TextAsset)']
    for i, v in enumerate(vals, 1):
        c = toc.cell(new_row, i, v)
        c.font = F_BODY
        c.alignment = WRAP
    end = max(last, new_row)
    toc.cell(5, 2, '=SUM(B7:B%d)' % end)
    toc.cell(5, 3, '=SUM(C7:C%d)' % end)
    print('Dictionary: %d mục = %d dòng (title %d, text %d), TOC row %d' % (
        len(data), row - 2, n_title, n_text, new_row))

    bak = xlsx + '.bak'
    shutil.copy2(xlsx, bak)
    wb.save(xlsx)
    print('đã lưu %s (backup %s)' % (xlsx, os.path.basename(bak)))


if __name__ == '__main__':
    main()
