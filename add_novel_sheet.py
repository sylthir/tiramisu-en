"""Thêm/cập nhật sheet truyện ngắn DLC từ bundle novel/novel_aoc01.
Mỗi dòng truyện = 1 dòng sheet, ID Novel|<tên>|<số dòng 0-based>. Dòng trống giữ nguyên.
Dùng: python add_novel_sheet.py [--xlsx ...]
"""
import argparse, json, os, re, shutil, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import UnityPy

HERE = os.path.dirname(os.path.abspath(__file__))
HEAD = ['ID', 'Tiếng Nhật / gốc', 'Bản dịch tiếng Anh', 'Loại text', 'Ngữ cảnh', 'Ghi chú của người dịch']
F_HEAD = Font(name='Arial', size=11, bold=True, color='FFFFFF')
F_BODY = Font(name='Arial', size=11, color='252525')
FILL_HEAD = PatternFill('solid', fgColor='314F4C')
FILL_TR = PatternFill('solid', fgColor='FFF3CF')
WRAP = Alignment(wrap_text=True, vertical='top')
BAD = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean(v):
    return BAD.sub('', v) if isinstance(v, str) else v


def main():
    cfg = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=cfg['workbook'])
    a = ap.parse_args()
    xlsx = os.path.join(HERE, a.xlsx) if not os.path.isabs(a.xlsx) else a.xlsx
    print('đang mở', xlsx)
    wb = openpyxl.load_workbook(xlsx)
    old = {}
    for dlc in cfg.get('dlc', []):
        sheet = dlc.get('sheet')
        if sheet in wb.sheetnames:
            for r in wb[sheet].iter_rows(min_row=2, values_only=True):
                if r and r[0] and r[2]:
                    old[str(r[0])] = str(r[2])

    for dlc in cfg.get('dlc', []):
        sheet = dlc.get('sheet') or 'DLC_Novel'
        romfs = os.path.expandvars(dlc['romfs_dir'].strip().strip('"'))
        bundle = os.path.join(romfs, *dlc['bundle'].replace('\\', '/').split('/'))
        if not os.path.isfile(bundle):
            print('! không thấy', bundle); continue
        env = UnityPy.load(bundle)
        rows = []
        for o in env.objects:
            if o.type.name != 'TextAsset':
                continue
            t = o.read()
            name = t.m_Name
            data = t.m_Script
            if isinstance(data, str):
                data = data.encode('utf-8')
            txt = data.decode('utf-8-sig')
            for idx, line in enumerate(txt.replace('\r\n', '\n').split('\n')):
                if not line.strip():
                    continue
                rid = 'Novel|%s|%d' % (name, idx)
                rows.append([rid, line, old.get(rid), 'Truyện ngắn DLC', '%s / dòng %d' % (name, idx + 1),
                             'Dịch truyện ngắn DLC (giữ dòng trống và thụt đầu dòng)'])
        if sheet in wb.sheetnames:
            wb.remove(wb[sheet])
        ws = wb.create_sheet(sheet)
        ws.freeze_panes = 'A2'
        for i, w in enumerate((26, 70, 70, 20, 24, 40), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for i, h in enumerate(HEAD, 1):
            c = ws.cell(1, i, h); c.font, c.fill = F_HEAD, FILL_HEAD
        for j, r in enumerate(rows, 2):
            for i, v in enumerate(r, 1):
                c = ws.cell(j, i, clean(v)); c.font, c.alignment = F_BODY, WRAP
                if i == 3: c.fill = FILL_TR
        if 'Dictionary' in wb.sheetnames:
            wb.move_sheet(ws, offset=wb.sheetnames.index('Dictionary') + 1 - wb.sheetnames.index(sheet))
        print('%s: %d dòng (đã dịch %d)' % (sheet, len(rows), sum(1 for r in rows if r[2])))

        toc = wb['TOC']
        last = 0
        for r in range(7, 400):
            v = toc.cell(r, 1).value
            if v not in (None, ''):
                last = r
            elif last and r > last + 2:
                break
        new_row = next((r for r in range(7, last + 1) if str(toc.cell(r, 1).value) == sheet), last + 1)
        vals = [sheet, len(rows), '=B%d-COUNTBLANK(%s!C2:C%d)' % (new_row, sheet, len(rows) + 1),
                '=C%d/B%d' % (new_row, new_row), 'Truyện ngắn DLC (%s)' % dlc['bundle']]
        for i, v in enumerate(vals, 1):
            c = toc.cell(new_row, i, v); c.font = F_BODY; c.alignment = WRAP
        end = max(last, new_row)
        toc.cell(5, 2, '=SUM(B7:B%d)' % end)
        toc.cell(5, 3, '=SUM(C7:C%d)' % end)

    bak = xlsx + '.bak'
    shutil.copy2(xlsx, bak)
    wb.save(xlsx)
    print('đã lưu', xlsx)


if __name__ == '__main__':
    main()
