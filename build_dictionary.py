"""Ghi từ điển tiếng Anh vào TextAsset 'dictionary' (và 'dictionary_en') trong resources.assets.
Đọc sheet 'Dictionary' (ID Dict|<no>|title / Dict|<no>|text). Ô C trống -> giữ nguyên tiếng Nhật.
Dùng: python build_dictionary.py            -> dist/resources.assets
"""
import argparse, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import gamedata as gd
from serialized import SF

try:
    import UnityPy
except ImportError:
    raise SystemExit('Cần UnityPy:  pip install UnityPy')


def read_sheet(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if 'Dictionary' not in wb.sheetnames:
        raise SystemExit('Workbook không có sheet Dictionary — chạy add_dictionary_sheet.py trước')
    ws = wb['Dictionary']
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    if not hdr or hdr[0] != 'ID':
        raise SystemExit('Sheet Dictionary thiếu cột ID')
    ci = 2
    for i, h in enumerate(hdr):
        if h in ('Bản dịch tiếng Anh', 'Tiếng Anh', 'Bản dịch'):
            ci = i
            break
    out = {}
    rx = re.compile(r'^Dict\|(\d+)\|(title|text)$')
    for row in it:
        if not row or not row[0]:
            continue
        m = rx.match(str(row[0]))
        if not m:
            continue
        no, kind = m.group(1), m.group(2)
        v = row[ci] if len(row) > ci else None
        if v is not None and str(v).strip():
            out.setdefault(no, {})[kind] = str(v).strip()
    return out


def main():
    cfg = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=cfg['workbook'])
    ap.add_argument('--base', default=os.path.join(HERE, 'dict', 'dictionary.json'),
                    help='dictionary.json gốc (thứ tự, category, ruby)')
    a = ap.parse_args()

    en = read_sheet(os.path.join(HERE, a.xlsx) if not os.path.isabs(a.xlsx) else a.xlsx)
    base = json.load(open(a.base, encoding='utf-8'))['data']
    entries = []
    n_en = 0
    for e in base:
        no = str(e.get('no', ''))
        t = e.get('title', '')
        x = e.get('text', '')
        tr = en.get(no, {})
        if tr.get('title'):
            t = tr['title']
            n_en += 1
        if tr.get('text'):
            x = tr['text']
        entries.append({'category': e.get('category', ''), 'no': no, 'title': t,
                        'text': x, 'ruby': e.get('ruby', '')})
    data = json.dumps({'data': entries}, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    print('Dictionary: %d mục, %d tên đã dịch, %d giải thích đã dịch' % (
        len(entries), n_en, sum(1 for e in base if en.get(str(e.get('no', '')), {}).get('text'))))

    romfs = os.path.expandvars(cfg['romfs_dir'].strip().strip('"'))
    src = os.path.join(romfs, 'Data', 'resources.assets')
    if not os.path.isfile(src):
        raise SystemExit('Không thấy %s' % src)
    env = UnityPy.load(src)
    names = {}
    for o in env.objects:
        if o.type.name == 'TextAsset':
            try:
                n = o.read().m_Name
            except Exception:
                continue
            if n in ('dictionary', 'dictionary_en'):
                names[n] = o.path_id
    if 'dictionary' not in names:
        raise SystemExit('resources.assets không có TextAsset dictionary')
    sf = SF(open(src, 'rb').read())
    repl = {}
    for n, pid in names.items():
        raw = sf.raw(pid)
        nm, p, _ = gd.ta_parse(raw)
        repl[pid] = gd.ta_make(raw, data)
    out = sf.rebuild(repl)
    out_dir = os.path.join(HERE, cfg.get('out_dir', 'dist'))
    os.makedirs(out_dir, exist_ok=True)
    dst = os.path.join(out_dir, 'resources.assets')
    open(dst, 'wb').write(out)

    # đọc lại kiểm tra
    env2 = UnityPy.load(dst)
    o = next(x for x in env2.objects if x.type.name == 'TextAsset' and x.read().m_Name == 'dictionary')
    nm, p, raw = gd.ta_parse(o.get_raw_data())
    back = json.loads(raw.decode('utf-8-sig'))
    ok = len(back.get('data', [])) == len(entries) and back['data'][0]['title'] == entries[0]['title']
    print('resources.assets -> %s (%.1f MB) | kiểm tra: %s' % (dst, len(out) / 1e6, 'OK' if ok else 'SAI'))
    if not ok:
        raise SystemExit('Đọc lại không khớp')


if __name__ == '__main__':
    main()
