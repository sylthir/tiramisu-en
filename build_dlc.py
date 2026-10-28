"""Build truyện ngắn DLC: ghi bản dịch sheet DLC_Novel_* vào bundle novel/novel_aoc0x.
Dùng: python build_dlc.py
Kết quả: dist/novel_aoc01 (build_dist sẽ đóng gói vào dist/<title_id>/<mod>/romfs/novel/).
"""
import argparse, json, os, re, shutil, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import gamedata as gd
from serialized import SF
from unityfs import UnityFS

try:
    import UnityPy
    import openpyxl
except ImportError as e:
    raise SystemExit('Cần UnityPy + openpyxl:  %s' % e)

RX = re.compile(r'^Novel\|(.+)\|(\d+)$')


def read_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        print('! workbook chưa có sheet %s — bỏ qua' % sheet)
        return {}
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    ci = 2
    for i, h in enumerate(hdr):
        if h in ('Bản dịch tiếng Anh', 'Tiếng Anh', 'Bản dịch'):
            ci = i
            break
    out = {}
    for row in it:
        if not row or not row[0]:
            continue
        m = RX.match(str(row[0]))
        if not m:
            continue
        v = row[ci] if len(row) > ci else None
        if v is not None and str(v).strip():
            out[(m.group(1), int(m.group(2)))] = str(v).strip()
    return out


def main():
    cfg = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=cfg['workbook'])
    a = ap.parse_args()
    xlsx = os.path.join(HERE, a.xlsx) if not os.path.isabs(a.xlsx) else a.xlsx
    out_dir = os.path.join(HERE, cfg.get('out_dir', 'dist'))
    os.makedirs(out_dir, exist_ok=True)

    for dlc in cfg.get('dlc', []):
        sheet = dlc.get('sheet', 'DLC_Novel')
        trans = read_sheet(xlsx, sheet)
        romfs = os.path.expandvars(dlc['romfs_dir'].strip().strip('"'))
        rel = dlc['bundle'].replace('\\', '/')
        src = os.path.join(romfs, *rel.split('/'))
        if not os.path.isfile(src):
            print('! không thấy', src); continue
        env = UnityPy.load(src)
        groups = {}
        for o in env.objects:
            if o.type.name != 'TextAsset':
                continue
            try:
                name = o.read().m_Name
            except Exception:
                continue
            groups.setdefault(o.assets_file.name, {})[o.path_id] = name
        u = UnityFS(src)
        datas = [u.node_data(i) for i in range(len(u.nodes))]
        patched = 0
        for ni, nd in enumerate(u.nodes):
            if nd[3] not in groups:
                continue
            sf = SF(datas[ni])
            repl = {}
            for pid, name in groups[nd[3]].items():
                raw = sf.raw(pid)
                nm, p, data = gd.ta_parse(raw)
                bom = data.startswith(b'\xef\xbb\xbf')
                txt = data.decode('utf-8-sig')
                lines = txt.replace('\r\n', '\n').split('\n')
                new = []
                for idx, line in enumerate(lines):
                    t = trans.get((name, idx))
                    if t:
                        t = t.replace('\r', ' ').replace('\n', ' ')
                        pref = re.match(r'^[\s　]*', line).group(0)
                        if pref and not re.match(r'^[\s　]', t):
                            t = pref + t
                        new.append(t); patched += 1
                    else:
                        new.append(line)
                ndata = '\r\n'.join(new).encode('utf-8')
                if bom:
                    ndata = b'\xef\xbb\xbf' + ndata
                repl[pid] = gd.ta_make(raw, ndata)
            datas[ni] = sf.rebuild(repl)
        dst = os.path.join(out_dir, os.path.basename(src))
        u.save(dst, datas)
        # xuất luôn vào cây mod theo title
        export = os.path.join(out_dir, dlc.get('title_id', ''), cfg.get('mod_name', 'English'),
                              'romfs', *rel.split('/'))
        os.makedirs(os.path.dirname(export), exist_ok=True)
        shutil.copy2(dst, export)
        print('  ->', export)
        # kiểm tra đọc lại
        e2 = UnityPy.load(dst)
        names = []
        for o in e2.objects:
            if o.type.name == 'TextAsset':
                names.append(o.read().m_Name)
        print('%s: %d dòng dịch, %d TextAsset (%s) -> %s (%.1f KB)' % (
            dlc.get('title_id', '?'), patched, len(names), ', '.join(sorted(names)), dst, os.path.getsize(dst) / 1e3))


if __name__ == '__main__':
    main()
