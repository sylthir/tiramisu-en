"""Dịch các ô còn trống (cột C) THEO LÔ nhiều dòng bằng Gemini / Groq / DeepL.
- Thẻ (<tips=..>, <param=..>, {0}...) được thay bằng placeholder ⟦n⟧ trước khi gửi và khôi phục sau khi dịch; thiếu/thừa placeholder thì dịch lại.
- Luật dịch lấy từ sheet Huong_dan + bảng xưng hô xung_ho.tsv.
- Lưu tiến độ từng dòng (progress_en.jsonl), ghi Excel định kỳ; hết hạn mức thì dừng, lần sau chạy lại dịch tiếp các ô còn trống.
Dùng:  python translate.py                        dịch tiếp mọi ô trống
       python translate.py --start Leon_01_01 sheet này làm trước
       python translate.py --sheet Common_00_01   chỉ các sheet này
       python translate.py --batch 10 --limit 50   lô 10 dòng, thử 50 dòng
       python translate.py --redo Leon_01_01 dịch lại cả sheet
       python translate.py --redo-flagged          dịch lại các ô tô đỏ
Khoá: GEMINI_API_KEY, GROQ_API_KEY, DEEPL_API_KEY (biến môi trường hoặc file .env)."""
import argparse, collections, json, os, re, sys, time, urllib.error, urllib.request
import openpyxl
from xungho import Table, fix_text, RUBY

HERE = os.path.dirname(os.path.abspath(__file__))
SKIP = ('TOC', 'Huong_dan')
PRIORITY = ('Character_names', 'Speaker_labels', 'Dic_terms', 'Dictionary')   # dịch trước để thuật ngữ nhất quán
PROGRESS = os.path.join(HERE, 'progress_en.jsonl')
UA = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36'
TAG = re.compile(r'</?[A-Za-z]+(?:=[^>]*)?>|\{\d+\}|\[[^\]]+\]')
PH = re.compile(r'⟦(\d+)⟧')
# Dấu hiệu prompt bị lẫn vào bản dịch (ví dụ khi model/API trả lại nguyên prompt)
PROMPT_LEAK = re.compile(r'GLOSSARY:|SPEAKER:|CONTEXT \(|TRANSLATE \(|NOTES PER ITEM|PREVIOUS LINES|^LINES:|speaker: |^\[\d+\] ', re.M)
SYMBOLS = re.compile(r'^[\s…・。、！？!?.,~～♪♡☆★ー\-—―（）()「」『』\u3000]*$')
SYM_MAP = str.maketrans({'…': '...', '・': '.', '。': '.', '、': ',', '！': '!', '？': '?', '～': '~', '\u3000': ' ',
                         '（': '(', '）': ')', '「': '"', '」': '"', '『': '“', '』': '”', 'ー': '-'})
SYSTEM = """You translate lines of a Japanese otome / mystery visual novel ("Le Mirage Mystique", set in 1910s Paris) into natural, fluent English for a game script.
RULES (follow strictly):
1. Output exactly one line per input item, in the form "[n] translation", with the same number n. No notes, no extra text.
2. Tokens like ⟦1⟧ stand for control tags ([主人公] = the player's name, <dic=..> glossary terms, <sprite ..>). Copy every token exactly once, in the same position. Never delete, change or invent them. Text between a <dic> pair is a glossary term: translate it, keep the tokens around it.
3. Dialogue 「」 -> "..." ; （） -> (...) ; 『』 -> “...” ; narration has no quotation marks. Each item is one line.
4. Names: use the glossary spelling exactly (Given Family order). Female characters (she/her): {females}. Characters with a name in the glossary who are not listed as female are male (he/him/his). For other unnamed characters infer gender from the Japanese and context; never default to male. Never use singular "they" for a named character.
5. In dialogue "I" is the speaker and "you" is the listener. Narration without a speaker is usually the heroine Aria speaking in first person (わたし).
6. {honor}
7. Project translation guide (follow it):
{guide}"""
HONOR_ON = 'Keep Japanese honorifics as suffixes (-san, -kun, -chan, -sama, -senpai). Do not use Mr./Ms.'
HONOR_OFF = 'Do not keep Japanese honorifics (-san, -kun, -sama, -senpai). Use first names, or Mr./Ms./Madame/Monsieur only where natural for the setting.'


class Daily(Exception): pass
class Rate(Exception):
    def __init__(self, wait): self.wait = wait
class Fatal(Exception): pass
class Transient(Exception): pass


def load_env():
    p = os.path.join(HERE, '.env')
    if os.path.exists(p):
        for line in open(p, encoding='utf-8'):
            if '=' in line and not line.lstrip().startswith('#'):
                k, v = line.strip().split('=', 1)
                if v.strip():
                    os.environ.setdefault(k.strip(), v.strip())


def wait_from(msg, hdr=None):
    if hdr:
        try:
            return float(hdr)
        except ValueError:
            pass
    m = re.search(r'(\d+)m(\d+(?:\.\d+)?)s', msg)
    if m:
        return int(m.group(1)) * 60 + float(m.group(2))
    m = re.search(r'retryDelay"?\s*:\s*"?(\d+(?:\.\d+)?)s', msg) or re.search(r'in (\d+(?:\.\d+)?)s', msg)
    return float(m.group(1)) if m else 30.0


def http_post(url, body, headers):
    headers = dict(headers, **{'User-Agent': UA, 'Accept': 'application/json', 'Content-Type': 'application/json'})
    req = urllib.request.Request(url, json.dumps(body).encode('utf-8'), headers, method='POST')
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            return json.loads(r.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        msg = e.read().decode('utf-8', 'replace')[:1500]
        if e.code == 456:                                  # DeepL: quota exceeded
            raise Daily(msg[:150])
        if e.code == 429:
            if re.search(r'PerDay|per day|requests per day|tokens per day|daily|quota', msg, re.I):
                raise Daily(msg[:150])
            raise Rate(wait_from(msg, e.headers.get('Retry-After')))
        if e.code in (400, 401, 403, 404):
            raise Fatal('HTTP %d: %s' % (e.code, msg[:200]))
        raise Transient('HTTP %d' % e.code)
    except (urllib.error.URLError, TimeoutError) as e:
        raise Transient(str(e))


def llm_call(slot, system, user):
    if slot['engine'] == 'gemini':
        r = http_post('https://generativelanguage.googleapis.com/v1beta/models/%s:generateContent' % slot['model'],
                      {'systemInstruction': {'parts': [{'text': system}]}, 'contents': [{'role': 'user', 'parts': [{'text': user}]}],
                       'generationConfig': {'temperature': 0.3}},
                      {'x-goog-api-key': os.environ['GEMINI_API_KEY']})
        return r['candidates'][0]['content']['parts'][0]['text']
    if slot['engine'] == 'groq':
        r = http_post('https://api.groq.com/openai/v1/chat/completions',
                      {'model': slot['model'], 'temperature': 0.3,
                       'messages': [{'role': 'system', 'content': system}, {'role': 'user', 'content': user}]},
                      {'Authorization': 'Bearer ' + os.environ['GROQ_API_KEY']})
        return r['choices'][0]['message']['content']
    if slot['engine'] == 'fake':                              # chỉ để thử luồng chạy
        slot['calls'] = slot.get('calls', 0) + 1
        if slot.get('quota') and slot['calls'] > slot['quota']:
            raise Daily('fake quota')
        body = user.split('LINES:')[-1]
        return '\n'.join('[%s] [EN] %s' % (m.group(1), m.group(2)) for m in re.finditer(r'^\[(\d+)\] (.*)$', body, re.M))
    raise Fatal('engine không hợp lệ: ' + slot['engine'])


def deepl_call(slot, texts, context):
    key = os.environ['DEEPL_API_KEY']
    url = 'https://api-free.deepl.com/v2/translate' if key.endswith(':fx') else 'https://api.deepl.com/v2/translate'
    body = {'text': texts, 'source_lang': 'JA', 'target_lang': 'EN-US', 'preserve_formatting': True}
    if context:
        body['context'] = context[:3000]
    r = http_post(url, body, {'Authorization': 'DeepL-Auth-Key ' + key})
    return [t['text'] for t in r['translations']]


def protect(plain):
    tags = []

    def rep(m):
        tags.append(m.group(0))
        return '⟦%d⟧' % len(tags)
    return TAG.sub(rep, plain), tags


def restore(text, tags):
    """-> (text with tags back, ok). ok=False when placeholders are missing, repeated or invented."""
    found = [int(x) for x in PH.findall(text)]
    if sorted(found) != list(range(1, len(tags) + 1)):
        return None, False
    return PH.sub(lambda m: tags[int(m.group(1)) - 1], text), True


def clean(t):
    t = re.sub(r'^```\w*\s*$|^```\s*$', '', t, flags=re.M)
    return t.strip()


def parse_numbered(text):
    out = {}
    for m in re.finditer(r'^\s*\[(\d+)\]\s*(.+?)\s*$', clean(text), re.M):
        out[int(m.group(1))] = m.group(2)
    return out


class Item:
    __slots__ = ('ws', 'row', 'ci', 'ni', 'ti', 'rid', 'jp', 'plain', 'typ', 'tags', 'prot', 'spk', 'who', 'prev', 'gloss', 'attempts')

    def __init__(self, **kw):
        for k, v in kw.items():
            setattr(self, k, v)
        self.attempts = 0


def item_context(ws, row, ci, tb, jp, typ):
    spk_jp = spk_en = None
    if ws.cell(row - 1, 1).value == 'speaker':
        spk_jp, spk_en = ws.cell(row - 1, 2).value, ws.cell(row - 1, ci).value
    prev, r = [], row - 1
    while r >= 2 and len(prev) < 3:
        rid = ws.cell(r, 1).value
        if rid and rid != 'speaker':
            j, e = ws.cell(r, 2).value, ws.cell(r, ci).value
            if j:
                prev.append((RUBY.sub(r'\1', str(j)).replace('\n', ' ')[:120], str(e or '').replace('\n', ' ')[:140]))
        r -= 1
    prev.reverse()
    info = tb.speaker(spk_jp) if spk_jp else None
    speaker = ('%s (%s)' % info) if info else (spk_en or '')
    return speaker, prev, spk_jp


def unify_glossary(tb, items):
    gl = {}
    for it in items:
        for jp, en, g in tb.relevant(it.plain, it.spk or ''):
            gl[en] = (jp, g)
    return '; '.join('%s = %s (%s)' % (en, jp, g) for en, (jp, g) in gl.items())


def load_guide(wb):
    if 'Huong_dan' not in wb.sheetnames:
        return ''
    lines = []
    for r in wb['Huong_dan'].iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[1]:
            lines.append('- %s: %s' % (r[0], r[1]))
    return '\n'.join(lines)[:2500]


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    load_env()
    cfg = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=cfg['workbook'])
    ap.add_argument('--sheet', nargs='*')
    ap.add_argument('--start', help='sheet dịch trước tiên')
    ap.add_argument('--redo', nargs='*', help='dịch lại cả các sheet này')
    ap.add_argument('--redo-flagged', action='store_true')
    ap.add_argument('--limit', type=int, default=0)
    ap.add_argument('--batch', type=int, default=0, help='số dòng mỗi lô (mặc định theo slot)')
    ap.add_argument('--engine', default='auto', help='auto | gemini | groq | deepl | fake')
    ap.add_argument('--save-every', type=int, default=100)
    a = ap.parse_args()

    keys = {'gemini': 'GEMINI_API_KEY', 'groq': 'GROQ_API_KEY', 'deepl': 'DEEPL_API_KEY'}
    if a.engine == 'fake':
        slots = [{'engine': 'fake', 'model': 'fake', 'interval': 0, 'batch': 10, 'quota': int(os.environ.get('FAKE_QUOTA', '0') or 0)}]
    else:
        slots = [dict(s) for s in cfg['slots'] if os.environ.get(keys.get(s['engine'], ''), '') and a.engine in ('auto', s['engine'])]
    if not slots:
        sys.exit('Chưa có khoá API nào (GEMINI_API_KEY / GROQ_API_KEY / DEEPL_API_KEY). Xem .env.example.')
    for s in slots:
        s.update(next_ok=0.0, dead=False, batch=a.batch or s.get('batch', 10))

    tb = Table()
    path = os.path.join(HERE, a.xlsx)
    print('đang mở', path, '...')
    wb = openpyxl.load_workbook(path)
    guide = load_guide(wb)
    dic_en = {}
    if 'Dic_terms' in wb.sheetnames:
        for r in wb['Dic_terms'].iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[2] and str(r[0]).startswith('Dic|'):
                dic_en[str(r[0])[4:]] = str(r[2])
    females = ', '.join(tb.females)
    done_prev = {}
    if os.path.exists(PROGRESS):
        for line in open(PROGRESS, encoding='utf-8'):
            try:
                d = json.loads(line)
                done_prev[d['id']] = d
            except ValueError:
                pass

    redo = set(a.redo or [])
    work, applied, filled_n, total_n, order = [], 0, {}, {}, {}
    for ws in wb.worksheets:
        if ws.title in SKIP:
            continue
        excluded = bool(a.sheet and ws.title not in a.sheet and ws.title not in redo)
        if excluded and not done_prev:
            continue
        hdr = [c.value for c in ws[1]]
        if not hdr or hdr[0] != 'ID':
            continue
        ci = next((i for i, h in enumerate(hdr) if h in ('Bản dịch tiếng Anh', 'Tiếng Anh', 'Bản dịch')), 2) + 1
        ti = hdr.index('Loại text') + 1 if 'Loại text' in hdr else 4
        ni = hdr.index('Ghi chú của người dịch') + 1 if 'Ghi chú của người dịch' in hdr else 6
        order[ws.title] = len(order)
        for row in range(2, ws.max_row + 1):
            rid, jp = ws.cell(row, 1).value, ws.cell(row, 2).value
            if not rid or rid == 'speaker' or not jp:
                continue
            en = ws.cell(row, ci).value
            blank = not (en and str(en).strip())
            total_n[ws.title] = total_n.get(ws.title, 0) + 1
            filled_n[ws.title] = filled_n.get(ws.title, 0) + (not blank)
            if str(rid) in done_prev and blank:
                ws.cell(row, ci).value = done_prev[str(rid)]['en']
                ws.cell(row, ni).value = done_prev[str(rid)].get('note') or ws.cell(row, ni).value
                applied += 1
                continue
            if excluded:
                continue
            flagged = a.redo_flagged and any(k in str(ws.cell(row, ni).value or '') for k in ('KT đại từ', 'Lỗi thẻ'))
            if blank or ws.title in redo or flagged:
                work.append((ws, row, ci, ni, ti, str(rid), str(jp), ws.cell(row, ti).value or ''))
    if applied:
        print('khôi phục %d dòng từ progress_en.jsonl' % applied)
    started = {s for s in filled_n if 0 < filled_n[s] < total_n[s]}
    work.sort(key=lambda w: (0 if w[0].title == a.start else 1, 0 if w[0].title in PRIORITY else 1, 0 if w[0].title in started else 1, order.get(w[0].title, 0)))
    if started:
        print('làm tiếp các sheet đang dở: ' + ', '.join(sorted(started, key=lambda s: order.get(s, 0))[:12]))
    if a.limit:
        work = work[:a.limit]

    # ---- chuẩn bị các dòng ----
    items, failed, local_done = [], [], 0
    for ws, row, ci, ni, ti, rid, jp, typ in work:
        plain = RUBY.sub(r'\1', jp)
        if SYMBOLS.match(TAG.sub('', plain)):
            en = plain.translate(SYM_MAP)
            ws.cell(row, ci).value = en
            ws.cell(row, ni).value = 'MT:local'
            local_done += 1
            continue
        speaker, prev, spk_jp = item_context(ws, row, ci, tb, jp, typ)
        prot, tags = protect(plain)
        prot = prot.replace('\n', '⏎')                      # xuống dòng gửi đi dưới dạng ⏎ (model không được tách dòng)
        items.append(Item(ws=ws, row=row, ci=ci, ni=ni, ti=ti, rid=rid, jp=jp, plain=plain, typ=typ, tags=tags, prot=prot,
                          spk=spk_jp, who=speaker, prev=prev, gloss=None))
    print('cần dịch %d dòng (%d ký hiệu xử lý tại chỗ) | slot: %s' % (
        len(items), local_done, ', '.join('%s/%s lô=%d' % (s['engine'], s['model'], s['batch']) for s in slots)))

    prog = open(PROGRESS, 'a', encoding='utf-8')
    st = dict(done=0, since=0)
    queue = collections.deque(items)

    def save(final=False):
        tmp = path + '.tmp'
        try:
            wb.save(tmp)
            os.replace(tmp, path)
        except PermissionError:
            print('  ! Không ghi được Excel (đang mở?). Tiến độ vẫn ở progress_en.jsonl, đóng Excel rồi chạy lại.')
            return
        st['since'] = 0
        print('  đã lưu Excel (%d dòng trong lần chạy này)' % st['done'])

    def apply(it, out, final_try=False):
        out = out.replace('⏎', '\n' if it.typ == 'Tên chương' else ' ')
        if PROMPT_LEAK.search(out) or len(out) > 4 * len(it.plain) + 80:
            return 'bản dịch lẫn prompt hoặc quá dài, bỏ qua'
        restored, ok = restore(out, it.tags)
        if not ok:
            if not final_try:
                return 'placeholder'
            # Lần cuối: giữ câu dịch, thẻ nào model làm mất thì bỏ và ghi chú (tô đỏ nhạt để kiểm tra)
            def keep(m):
                k = int(m.group(1))
                return it.tags[k - 1] if 1 <= k <= len(it.tags) else ''
            restored = PH.sub(keep, out)
        final, notes = fix_text(tb, it.jp, restored, it.typ)
        return final, notes

    stop = None
    try:
        while queue:
            alive = [s for s in slots if not s['dead']]
            if not alive:
                stop = 'Hết hạn mức ở mọi slot.'
                break
            now = time.time()
            ready = [s for s in alive if s['next_ok'] <= now]
            slot = ready[0] if ready else min(alive, key=lambda s: s['next_ok'])
            if not ready:
                time.sleep(max(0.2, slot['next_ok'] - now))
                continue
            batch = [queue.popleft() for _ in range(min(slot['batch'], len(queue)))]
            try:
                if slot['engine'] == 'deepl':
                    ctx = 'Japanese visual novel script. ' + guide
                    outs = deepl_call(slot, [it.prot for it in batch], ctx)
                    got = {i + 1: t for i, t in enumerate(outs)}
                else:
                    group = [it for it in batch]
                    gl = unify_glossary(tb, group)
                    system = SYSTEM.format(females=females, honor=HONOR_ON if cfg.get('honorific_rule', True) else HONOR_OFF,
                                           guide=guide or '(none)')
                    lines = ['[%d] %s' % (n, it.prot) for n, it in enumerate(group, 1)]
                    hints = []
                    for n, it in enumerate(group, 1):
                        if it.who:
                            hints.append('[%d] speaker: %s' % (n, it.who))
                        if it.typ and it.typ != 'Thoại':
                            hints.append('[%d] type: %s' % (n, it.typ))
                    for n, it in enumerate(group, 1):
                        terms = ['%s=%s' % (k, dic_en[k]) for k in re.findall(r'<dic=(\d+)>', it.plain) if k in dic_en]
                        if terms:
                            hints.append('[%d] glossary terms (use these English words inside <dic> tags): %s' % (n, '; '.join(terms)))
                    ctx = '\n'.join('  %s => %s' % p for p in group[0].prev) or '  (none)'
                    user = ('GLOSSARY: %s\nNOTES PER ITEM (do not translate these notes):\n%s\nPREVIOUS LINES (context only, do not translate):\n%s\nLINES:\n%s'
                            % (gl or '(none)', '\n'.join(hints) or '(none)', ctx, '\n'.join(lines)))
                    got = parse_numbered(llm_call(slot, system, user))
                    if not got and len(batch) > 1:          # đáp án không đúng định dạng: chia nhỏ lô
                        slot['batch'] = max(1, len(batch) // 2)
                        for it in reversed(batch):
                            queue.appendleft(it)
                        print('  %s trả về sai định dạng, giảm lô còn %d' % (slot['model'], slot['batch']))
                        continue
                slot['next_ok'] = time.time() + slot['interval']
            except Daily as ex:
                slot['dead'] = True
                for it in reversed(batch):
                    queue.appendleft(it)
                print('  slot %s/%s hết hạn mức (%s)' % (slot['engine'], slot['model'], str(ex)[:70]))
                continue
            except Rate as ex:
                wait = min(ex.wait, 120)
                slot['next_ok'] = time.time() + wait
                for it in reversed(batch):
                    queue.appendleft(it)
                print('  slot %s/%s giới hạn tốc độ, chờ %ds' % (slot['engine'], slot['model'], wait))
                continue
            except Fatal as ex:
                slot['dead'] = True
                for it in reversed(batch):
                    queue.appendleft(it)
                print('  slot %s/%s bị loại: %s' % (slot['engine'], slot['model'], str(ex)[:120]))
                if '1010' in str(ex):
                    print('  (1010 = Cloudflare chặn; tắt VPN/proxy)')
                continue
            except Transient as ex:
                slot['next_ok'] = time.time() + 10
                for it in reversed(batch):
                    queue.appendleft(it)
                print('  lỗi mạng/máy chủ (%s), thử lại' % str(ex)[:60])
                continue

            for n, it in enumerate(batch, 1):
                out = got.get(n)
                res = None if out is None else apply(it, out, final_try=it.attempts >= 2)
                if res is None or isinstance(res, str):
                    it.attempts += 1
                    if it.attempts >= 3:
                        failed.append('%s: %s' % (it.rid, 'không có kết quả' if res is None else res))
                    else:
                        queue.appendleft(it)
                    continue
                final, notes = res
                tag = 'MT:%s/%s' % (slot['engine'], slot['model'])
                old = [p for p in str(it.ws.cell(it.row, it.ni).value or '').split(' | ')
                       if p and not p.startswith(('MT:', 'Auto:', 'KT đại từ', 'Lỗi thẻ'))]
                note = ' | '.join([tag] + notes + old)
                it.ws.cell(it.row, it.ci).value = final
                it.ws.cell(it.row, it.ni).value = note
                prog.write(json.dumps({'id': it.rid, 'en': final, 'note': note}, ensure_ascii=False) + '\n')
                prog.flush()
                st['done'] += 1
                st['since'] += 1
            print('[%d còn lại] %-26s lô %d dòng' % (len(queue), slot['model'], len(batch)))
            if st['since'] >= a.save_every:
                save()
    except KeyboardInterrupt:
        stop = 'Đã dừng bằng Ctrl+C.'
    finally:
        save()
        prog.close()
    if failed:
        os.makedirs(os.path.join(HERE, 'logs'), exist_ok=True)
        open(os.path.join(HERE, 'logs', 'translate_failed.txt'), 'w', encoding='utf-8').write('\n'.join(failed))
        print('%d dòng bị bỏ qua (logs/translate_failed.txt), lần chạy sau sẽ thử lại' % len(failed))
    left = len(queue)
    print('\nxong lần chạy: dịch %d dòng, còn %d dòng chưa dịch.' % (st['done'], left))
    if stop:
        print(stop + ' Có thêm hạn mức thì chạy lại TRANSLATE.bat để dịch tiếp.')
        sys.exit(2)


if __name__ == '__main__':
    main()
