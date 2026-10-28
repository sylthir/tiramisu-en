"""Bảng xưng hô (tên + giới tính) dùng chung cho translate.py và fix_workbook.py:
chuẩn hoá tên theo xung_ho.tsv, dọn dòng thừa, kiểm tra / sửa đại từ he-she."""
import json
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
# Tên máy dịch hay viết sai -> tên đúng (bổ sung thêm trong config.json -> name_variants)
VARIANTS = [('Leon Vitri', 'Leon Vitry'), ('Léon', 'Leon'), ('Ulysses', 'Ulysse'), ('Ulysse Blanchard', 'Ulysse Blanchard'),
            ('Ahes', 'Ahès'), ('Aes', 'Ahès'), ('Aès', 'Ahès'), ('Rudi', 'Rudy'), ('Fin', 'Finn'), ('Charle', 'Charles'),
            ('Mavis Redmund', 'Mavis Redmond'), ('Serge Grac', 'Serge Gracq')]
MALE = re.compile(r'\b(he|him|his|himself)\b', re.I)
FEM = re.compile(r'\b(she|her|hers|herself)\b', re.I)
JP_M = re.compile(r'彼(?!女|ら)')
JP_F = re.compile(r'彼女(?!たち)')
JP_WOMAN = re.compile(r'女|姉|妹|母|娘|奥さん|お嬢|ママ|婦人|おばさん|嫁|祖母|夫人')
EN_WOMAN = re.compile(r'\b(woman|women|girl|girls|lady|ladies|mother|mom|sister|wife|daughter|female|madam|aunt|grandmother|madame)\b', re.I)
FUNC = set('the a an to at on in for with and but or that if so as about up down out from into because when while than then of by over after before off just not'.split())
RUBY = re.compile(r'<ruby=[^>]*>(.*?)</ruby>', re.S)


def _variants():
    try:
        cfg = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
        return list(cfg.get('name_variants', {}).items())
    except Exception:
        return []


class Table:
    def __init__(self, path=None):
        self.rows, self.tok = [], {}
        for line in open(path or os.path.join(HERE, 'xung_ho.tsv'), encoding='utf-8-sig'):
            if not line.strip() or line.startswith('#'):
                continue
            p = line.rstrip('\r\n').split('\t')
            if len(p) < 3:
                continue
            jp, en, g = re.sub(r'\s+', '', p[0]), p[1], ('F' if p[2].strip().lower().startswith('f') else 'M')
            self.rows.append((jp, en, g))
            for t in [t for t in re.split(r'[ ・]', en) if t]:
                self.tok[t] = g
        self.females = sorted({t for t, g in self.tok.items() if g == 'F'})
        self.variants = [v for v in VARIANTS if v[0] != v[1]] + _variants()

    # ---- prompt helpers ------------------------------------------------------
    def relevant(self, *jp_texts):
        norm = re.sub(r'\s+', '', ' '.join(t or '' for t in jp_texts))
        seen, out = set(), []
        for jp, en, g in sorted(self.rows, key=lambda r: -len(r[0])):
            if jp in norm and en not in seen and '・' not in jp:
                seen.add(en)
                out.append((jp, en, 'female' if g == 'F' else 'male'))
        return out

    def speaker(self, spk_cell):
        """(english name, gender word) cho ô người nói (vd 'レオン', '【ゼノ/？？？】') hoặc None."""
        if not spk_cell:
            return None
        key = re.sub(r'\s+', '', str(spk_cell)).strip('【】')
        first = key.split('/')[0]
        for jp, en, g in self.rows:
            if jp == first:
                return en, 'female' if g == 'F' else 'male'
        return None

    def label_en(self, label):
        """Dịch nhãn người nói khi mọi phần đều có trong bảng hoặc là ？？？ (vd 'ゼノ/？？？' -> 'Zeno/???'). Không đủ -> ''."""
        if label == 'player':
            return 'player'
        parts, out = label.split('/'), []
        for p in parts:
            p2 = re.sub(r'\s+', '', p)
            if re.fullmatch(r'[？?]+', p2):
                out.append('?' * len(p2))
                continue
            hit = next((en for jp, en, g in self.rows if jp == p2 and ' ' not in en), None)
            if not hit:
                return ''
            out.append(hit)
        return '/'.join(out)

    # ---- text fixing ---------------------------------------------------------
    def canon_names(self, t):
        for a, b in self.variants:
            t = re.sub(r'(?<![A-Za-z])%s(?![A-Za-z])' % re.escape(a), b, t)

        def title(m):
            g = self.tok.get(m.group(2))
            if g == 'M' and m.group(1) not in ('Mr', 'Monsieur'):
                return 'Mr. ' + m.group(2)
            if g == 'F' and m.group(1) == 'Mr':
                return 'Ms. ' + m.group(2)
            return m.group(0)
        return re.sub(r'\b(Mr|Ms|Mrs|Miss)\.?\s+([A-Z][A-Za-zè]+)', title, t)

    def names_in(self, s):
        return {w for w in re.findall(r"[A-Za-zè]+", s) if w in self.tok}


def tidy(t, keep_newline=True):
    """Bỏ thẻ ruby, dòng trống thừa, khoảng trắng thừa."""
    t = t.replace('\r\n', '\n').replace('\r', '\n')
    t = re.sub(r'(?<=[A-Za-z])</ruby>(?=[A-Za-z])', '</ruby> ', t)
    t = RUBY.sub(lambda m: m.group(1), t)
    lines = [re.sub(r'[ \t\u3000]+', ' ', ln).strip() for ln in t.split('\n')]
    lines = [ln for ln in lines if ln]
    return ('\n' if keep_newline else ' ').join(lines)


def _case(src, dst):
    return dst.capitalize() if src[:1].isupper() else dst


def swap_pronouns(s, to):
    """Đổi đại từ trong một câu sang giới tính `to` ('M' hoặc 'F')."""
    if to == 'M':
        def f(m):
            w = m.group(0); lw = w.lower()
            if lw == 'she':
                return _case(w, 'he')
            if lw == 'herself':
                return _case(w, 'himself')
            if lw == 'hers':
                return _case(w, 'his')
            nxt = re.match(r"\s*([A-Za-z'’]+)?", s[m.end():]).group(1) or ''
            end = s[m.end():].lstrip()[:1] in ('', '.', ',', '!', '?', ';', ':', '"', '”', ')', '…', '-')
            obj = end or nxt.lower() in FUNC or nxt.lower().endswith('ly')
            return _case(w, 'him' if obj else 'his')
        return re.sub(r'\b(she|her|hers|herself)\b', f, s, flags=re.I)
    m2 = {'he': 'she', 'him': 'her', 'his': 'her', 'himself': 'herself'}
    return re.sub(r'\b(he|him|his|himself)\b', lambda m: _case(m.group(0), m2[m.group(0).lower()]), s, flags=re.I)


def sentences(t):
    return [x for x in re.split(r'(?<=[.!?…])\s+|\n+', t) if x.strip()]


def pronoun_pass(tb, jp, en):
    """-> (new_en, fixed_count, unresolved_count).
    - Ô chỉ có nhân vật nam theo bảng + dùng she/her (không có he/him, không có dấu hiệu nữ trong JP/EN) -> đổi thành he/him.
    - Ô có nhân vật nữ hoặc có cả he lẫn she -> không tự đổi, đánh dấu KT để người dịch xem.
    - Ô có nhiều nhân vật khác giới -> xét từng câu có đúng một giới tính."""
    fixed = unresolved = 0
    out = en
    woman_ctx = bool(JP_WOMAN.search(jp or '')) or bool(EN_WOMAN.search(en))
    has_f, has_m = bool(FEM.search(en)), bool(MALE.search(en))
    if not (has_f or has_m):
        return out, 0, 0
    gs = {tb.tok[w] for w in tb.names_in(en)}
    if gs == {'M'}:
        if has_f and not has_m and not woman_ctx:
            out = swap_pronouns(en, 'M')
            fixed += len(FEM.findall(en))
        elif has_f:
            unresolved += 1
    elif gs == {'F'}:
        if re.search(r'\bhe\b', en, re.I):
            unresolved += 1
    elif len(gs) > 1:
        for s in sentences(en):
            sg = {tb.tok[w] for w in tb.names_in(s)}
            if len(sg) != 1:
                continue
            g = sg.pop()
            wrong = (FEM if g == 'M' else MALE).search(s)
            right = (MALE if g == 'M' else FEM).search(s)
            if not wrong:
                continue
            if right or g == 'F':
                unresolved += 1
                continue
            out = out.replace(s, swap_pronouns(s, 'M'), 1)
            fixed += 1
    if jp and not woman_ctx and gs == {'M'} and out == en:
        jm, jf = bool(JP_M.search(jp)), bool(JP_F.search(jp))
        if jm and not jf and has_f and not has_m:
            out, fixed = swap_pronouns(en, 'M'), fixed + 1
    return out, fixed, unresolved


def fix_text(tb, jp, en, typ='', keep_newline=True):
    """Làm sạch + chuẩn hoá tên + sửa đại từ. -> (text, notes list)."""
    notes = []
    before = tidy(en, keep_newline)
    new = tb.canon_names(before)
    new, fixed, unresolved = pronoun_pass(tb, jp or '', new)
    if fixed:
        notes.append('Auto: đổi đại từ he/she theo bảng xưng hô (%d chỗ)' % fixed)
    if unresolved:
        notes.append('KT đại từ: có thể sai he/she, hãy kiểm tra')
    return new, notes


def fit_widths(header, rows, min_w=8, max_w=60):
    n = len(header)
    widths = [len(str(h or '')) for h in header]
    for r in rows:
        for i in range(min(n, len(r))):
            v = r[i]
            if v is None or v == '':
                continue
            longest = max(len(x) for x in str(v).split('\n'))
            if longest > widths[i]:
                widths[i] = longest
    return [min(max(w + 2, min_w), max_w) for w in widths]


def fit_sheet(ws, max_w=60):
    from openpyxl.utils import get_column_letter
    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    widths = fit_widths(header, rows, max_w=max_w)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
